"""
Генератор шаблона Excel для массовой загрузки трек-кодов
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def generate_import_template(language: str = "ru") -> str:
    """
    Генерация шаблона Excel для импорта трек-кодов

    Args:
        language: Язык шаблона (ru или tj)

    Returns:
        Путь к созданному файлу
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Импорт трек-кодов" if language == "ru" else "Воридоти рамзҳо"

    # Заголовки
    headers = {
        "ru": [
            "Трек-код*",
            "Название товара",
            "Категория",
            "Количество",
            "Цена USD",
            "Вес кг",
            "Страна отправления"
        ],
        "tj": [
            "Рамзи тамошобин*",
            "Номи маҳсулот",
            "Гурӯҳ",
            "Миқдор",
            "Нарх USD",
            "Вазн кг",
            "Кишвари фиристод"
        ]
    }

    # Примеры данных
    examples = {
        "ru": [
            ["TRACK001", "Смартфон Samsung", "электроника", 1, 299.99, 0.5, "China"],
            ["TRACK002", "Кроссовки Nike", "обувь", 2, 89.50, 1.2, "China"],
            ["TRACK003", "Игрушка", "игрушки", 5, 15.00, 0.3, "China"]
        ],
        "tj": [
            ["TRACK001", "Смартфони Samsung", "электроника", 1, 299.99, 0.5, "Чин"],
            ["TRACK002", "Кроссовкаҳои Nike", "обувь", 2, 89.50, 1.2, "Чин"],
            ["TRACK003", "Бозича", "игрушки", 5, 15.00, 0.3, "Чин"]
        ]
    }

    # Добавляем заголовки
    ws.append(headers[language])

    # Стили для заголовков
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    alignment_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Применяем стили к заголовкам
    for col in range(1, len(headers[language]) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment_center
        cell.border = thin_border

    # Добавляем примеры
    for row_data in examples[language]:
        ws.append(row_data)

    # Применяем границы к примерам
    for row in range(2, 2 + len(examples[language])):
        for col in range(1, len(headers[language]) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            if col in [4, 5, 6]:  # Выравниваем числовые поля
                cell.alignment = Alignment(horizontal="right", vertical="center")

    # Настройка ширины колонок
    column_widths = [20, 30, 20, 12, 12, 12, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Высота заголовка
    ws.row_dimensions[1].height = 30

    # Добавляем инструкции на отдельном листе
    ws_instructions = wb.create_sheet("Инструкция" if language == "ru" else "Дастур")

    instructions = {
        "ru": [
            ["📋 ИНСТРУКЦИЯ ПО ИМПОРТУ ТРЕК-КОДОВ"],
            [""],
            ["1. ОБЯЗАТЕЛЬНЫЕ ПОЛЯ:"],
            ["   • Трек-код (колонка A) - обязательно для заполнения"],
            [""],
            ["2. ОПЦИОНАЛЬНЫЕ ПОЛЯ:"],
            ["   • Название товара (колонка B)"],
            ["   • Категория (колонка C)"],
            ["   • Количество (колонка D)"],
            ["   • Цена в долларах США (колонка E)"],
            ["   • Вес в килограммах (колонка F)"],
            ["   • Страна отправления (колонка G)"],
            [""],
            ["3. ДОПУСТИМЫЕ КАТЕГОРИИ:"],
            ["   • электроника"],
            ["   • одежда"],
            ["   • обувь"],
            ["   • бытовая техника"],
            ["   • косметика"],
            ["   • игрушки"],
            ["   • автозапчасти"],
            ["   • спорттовары"],
            ["   • другое"],
            [""],
            ["4. ТРЕБОВАНИЯ К ТРЕК-КОДУ:"],
            ["   • Минимум 6 символов"],
            ["   • Максимум 50 символов"],
            ["   • Допускаются: буквы (A-Z), цифры (0-9), дефис (-)"],
            ["   • Регистр не важен (автоматически преобразуется в верхний)"],
            [""],
            ["5. ВАЖНО:"],
            ["   • Дубликаты трек-кодов будут пропущены"],
            ["   • Первая строка (заголовки) не импортируется"],
            ["   • Пустые строки пропускаются"],
            ["   • Неверные данные будут отображены в отчёте об ошибках"],
            [""],
            ["6. ПОСЛЕ ИМПОРТА:"],
            ["   • Вы получите подробный отчёт о результатах"],
            ["   • Успешно импортированные трек-коды"],
            ["   • Пропущенные дубликаты"],
            ["   • Ошибки с описанием проблем"]
        ],
        "tj": [
            ["📋 ДАСТУРИ ВОРИДОТИ РАМЗҲОИ ТАМОШОБИН"],
            [""],
            ["1. МАЙДОНҲОИ ҲАТМӢ:"],
            ["   • Рамзи тамошобин (сутуни A) - ҳатман пур карда шавад"],
            [""],
            ["2. МАЙДОНҲОИ ИЛОВАГӢ:"],
            ["   • Номи маҳсулот (сутуни B)"],
            ["   • Гурӯҳ (сутуни C)"],
            ["   • Миқдор (сутуни D)"],
            ["   • Нарх дар доллари ИМА (сутуни E)"],
            ["   • Вазн дар килограм (сутуни F)"],
            ["   • Кишвари фиристод (сутуни G)"],
            [""],
            ["3. ГУРӮҲҲОИ ИҶОЗАТДОДАШУДА:"],
            ["   • электроника"],
            ["   • одежда (либос)"],
            ["   • обувь (пойафзол)"],
            ["   • бытовая техника (асбобҳои хонагӣ)"],
            ["   • косметика"],
            ["   • игрушки (бозичаҳо)"],
            ["   • автозапчасти"],
            ["   • спорттовары (ашёҳои варзишӣ)"],
            ["   • другое (дигар)"],
            [""],
            ["4. ТАЛАБОТҲО БА РАМЗИ ТАМОШОБИН:"],
            ["   • Ҳадди ақал 6 аломат"],
            ["   • Ҳадди аксар 50 аломат"],
            ["   • Иҷозат дода мешаванд: ҳарфҳо (A-Z), рақамҳо (0-9), дефис (-)"],
            [""],
            ["5. МУҲИМ:"],
            ["   • Дубликатҳои рамзҳо гузаронида мешаванд"],
            ["   • Сатри аввал (сарлавҳаҳо) ворид намешаванд"],
            ["   • Сатрҳои холӣ гузаронида мешаванд"],
            [""],
            ["6. БАЪД АЗ ВОРИД:"],
            ["   • Шумо гузориши муфассал дар бораи натиҷаҳо мегиред"]
        ]
    }

    for row_data in instructions[language]:
        ws_instructions.append(row_data)

    # Стили для инструкций
    title_font = Font(bold=True, size=14, color="366092")
    ws_instructions['A1'].font = title_font
    ws_instructions['A1'].alignment = Alignment(horizontal="left", vertical="center")

    # Ширина колонки
    ws_instructions.column_dimensions['A'].width = 80

    # Создаём папку для шаблонов, если её нет
    os.makedirs("templates", exist_ok=True)

    # Сохраняем файл
    filename = f"templates/import_template_{language}.xlsx"
    wb.save(filename)

    return filename


if __name__ == "__main__":
    # Генерируем шаблоны на русском и таджикском
    template_ru = generate_import_template("ru")
    print(f"✅ Создан русский шаблон: {template_ru}")

    template_tj = generate_import_template("tj")
    print(f"✅ Создан таджикский шаблон: {template_tj}")
