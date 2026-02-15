"""
Сервис массовой загрузки трек-кодов товаров
Поддерживает импорт из текстового списка и Excel файлов
"""
import logging
from typing import List, Dict, Tuple, Optional
from datetime import datetime
import re
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from database.models import Product, ProductStatus, ProductCategory
from database.repository import ProductRepository
from services.track_code_generator import TrackCodeGenerator

logger = logging.getLogger(__name__)


class BulkImportResult:
    """Результат массового импорта"""

    def __init__(self):
        self.success_count = 0
        self.duplicate_count = 0
        self.error_count = 0
        self.success_codes: List[str] = []
        self.duplicate_codes: List[str] = []
        self.errors: List[Dict[str, str]] = []

    def add_success(self, track_code: str):
        """Добавить успешно импортированный трек-код"""
        self.success_count += 1
        self.success_codes.append(track_code)

    def add_duplicate(self, track_code: str):
        """Добавить дубликат"""
        self.duplicate_count += 1
        self.duplicate_codes.append(track_code)

    def add_error(self, track_code: str, error: str):
        """Добавить ошибку"""
        self.error_count += 1
        self.errors.append({"track_code": track_code, "error": error})

    def get_summary(self, language: str = "ru") -> str:
        """Получить итоговый отчёт"""
        if language == "ru":
            summary = f"📊 Результаты массовой загрузки:\n\n"
            summary += f"✅ Успешно добавлено: {self.success_count}\n"
            summary += f"⚠️ Дубликатов пропущено: {self.duplicate_count}\n"
            summary += f"❌ Ошибок: {self.error_count}\n\n"

            if self.success_codes:
                summary += f"✅ Добавленные трек-коды:\n"
                for code in self.success_codes[:10]:  # Показываем первые 10
                    summary += f"  • {code}\n"
                if len(self.success_codes) > 10:
                    summary += f"  ... и ещё {len(self.success_codes) - 10}\n"
                summary += "\n"

            if self.duplicate_codes:
                summary += f"⚠️ Дубликаты:\n"
                for code in self.duplicate_codes[:5]:
                    summary += f"  • {code}\n"
                if len(self.duplicate_codes) > 5:
                    summary += f"  ... и ещё {len(self.duplicate_codes) - 5}\n"
                summary += "\n"

            if self.errors:
                summary += f"❌ Ошибки:\n"
                for err in self.errors[:5]:
                    summary += f"  • {err['track_code']}: {err['error']}\n"
                if len(self.errors) > 5:
                    summary += f"  ... и ещё {len(self.errors) - 5}\n"
        else:  # tj
            summary = f"📊 Натиҷаҳои боркунии оммавӣ:\n\n"
            summary += f"✅ Бомуваффақият илова шуд: {self.success_count}\n"
            summary += f"⚠️ Дубликатҳо гузаронида шуданд: {self.duplicate_count}\n"
            summary += f"❌ Хатогиҳо: {self.error_count}\n\n"

            if self.success_codes:
                summary += f"✅ Рамзҳои иловашуда:\n"
                for code in self.success_codes[:10]:
                    summary += f"  • {code}\n"
                if len(self.success_codes) > 10:
                    summary += f"  ... ва боз {len(self.success_codes) - 10}\n"
                summary += "\n"

            if self.duplicate_codes:
                summary += f"⚠️ Дубликатҳо:\n"
                for code in self.duplicate_codes[:5]:
                    summary += f"  • {code}\n"
                if len(self.duplicate_codes) > 5:
                    summary += f"  ... ва боз {len(self.duplicate_codes) - 5}\n"
                summary += "\n"

            if self.errors:
                summary += f"❌ Хатогиҳо:\n"
                for err in self.errors[:5]:
                    summary += f"  • {err['track_code']}: {err['error']}\n"
                if len(self.errors) > 5:
                    summary += f"  ... ва боз {len(self.errors) - 5}\n"

        return summary


class TrackCodeValidator:
    """Валидатор трек-кодов"""

    # Регулярное выражение для базовой валидации трек-кода
    # Минимум 6 символов, допускаются буквы, цифры, дефисы
    TRACK_CODE_PATTERN = re.compile(r'^[A-Z0-9\-]{6,50}$', re.IGNORECASE)

    @staticmethod
    def validate_track_code(track_code: str) -> Tuple[bool, Optional[str]]:
        """
        Валидация трек-кода

        Returns:
            (is_valid, error_message)
        """
        if not track_code or not track_code.strip():
            return False, "Пустой трек-код"

        track_code = track_code.strip().upper()

        if len(track_code) < 6:
            return False, "Слишком короткий (минимум 6 символов)"

        if len(track_code) > 50:
            return False, "Слишком длинный (максимум 50 символов)"

        if not TrackCodeValidator.TRACK_CODE_PATTERN.match(track_code):
            return False, "Недопустимые символы (разрешены: A-Z, 0-9, -)"

        return True, None

    @staticmethod
    def parse_track_codes_from_text(text: str) -> List[str]:
        """
        Извлечение трек-кодов из текста
        Поддерживает разделители: запятая, точка с запятой, новая строка, пробел
        """
        # Разделители
        delimiters = [',', ';', '\n', ' ', '\t']

        # Заменяем все разделители на запятую
        for delimiter in delimiters:
            text = text.replace(delimiter, ',')

        # Разбиваем по запятой и очищаем
        codes = [code.strip().upper() for code in text.split(',') if code.strip()]

        # Убираем дубликаты, сохраняя порядок
        seen = set()
        unique_codes = []
        for code in codes:
            if code not in seen:
                seen.add(code)
                unique_codes.append(code)

        return unique_codes


class BulkImportService:
    """Сервис массового импорта трек-кодов"""

    def __init__(self, product_repo: ProductRepository):
        self.product_repo = product_repo
        self.validator = TrackCodeValidator()

    async def import_from_text(
        self,
        text: str,
        user_id: int,
        default_category: str = "other",
        default_country: str = "China"
    ) -> BulkImportResult:
        """
        Массовый импорт из текстового списка

        Args:
            text: Текст со списком трек-кодов
            user_id: ID пользователя (админа)
            default_category: Категория по умолчанию
            default_country: Страна отправления по умолчанию

        Returns:
            BulkImportResult с результатами импорта
        """
        result = BulkImportResult()

        # Извлекаем трек-коды
        track_codes = self.validator.parse_track_codes_from_text(text)

        if not track_codes:
            logger.warning("Не найдено ни одного трек-кода в тексте")
            return result

        logger.info(f"Найдено {len(track_codes)} уникальных трек-кодов для импорта")

        # Обрабатываем каждый трек-код
        for track_code in track_codes:
            # Валидация
            is_valid, error_msg = self.validator.validate_track_code(track_code)
            if not is_valid:
                result.add_error(track_code, error_msg)
                continue

            # Проверка на дубликат
            existing = await self.product_repo.get_product_by_track_code(track_code)
            if existing:
                result.add_duplicate(track_code)
                continue

            # Создаём товар с минимальными данными
            try:
                await self.product_repo.create_product(
                    track_code=track_code,
                    user_id=user_id,
                    product_name=f"Товар {track_code}",
                    product_category=default_category,
                    country_from=default_country,
                    quantity=1,
                    unit_price_usd=0.0,
                    total_value_usd=0.0,
                    weight_kg=0.0,
                    status=ProductStatus.CREATED
                )
                result.add_success(track_code)
                logger.info(f"Успешно импортирован трек-код: {track_code}")
            except Exception as e:
                error_msg = f"Ошибка БД: {str(e)}"
                result.add_error(track_code, error_msg)
                logger.error(f"Ошибка при создании товара {track_code}: {e}")

        return result

    async def import_from_excel(
        self,
        file_path: str,
        user_id: int
    ) -> BulkImportResult:
        """
        Массовый импорт из Excel файла

        Ожидаемая структура файла:
        - Колонка A: Трек-код (обязательно)
        - Колонка B: Название товара (опционально)
        - Колонка C: Категория (опционально)
        - Колонка D: Количество (опционально)
        - Колонка E: Цена (опционально)
        - Колонка F: Вес (опционально)
        - Колонка G: Страна отправления (опционально)

        Args:
            file_path: Путь к Excel файлу
            user_id: ID пользователя (админа)

        Returns:
            BulkImportResult с результатами импорта
        """
        result = BulkImportResult()

        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
        except InvalidFileException:
            logger.error(f"Невалидный Excel файл: {file_path}")
            result.add_error("FILE", "Невалидный Excel файл")
            return result
        except Exception as e:
            logger.error(f"Ошибка открытия файла {file_path}: {e}")
            result.add_error("FILE", f"Ошибка открытия файла: {str(e)}")
            return result

        # Пропускаем заголовок (первая строка)
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        logger.info(f"Найдено {len(rows)} строк в Excel файле")

        # Маппинг категорий
        category_map = {
            "электроника": "electronics",
            "одежда": "clothing",
            "обувь": "shoes",
            "бытовая техника": "home_appliances",
            "косметика": "beauty",
            "игрушки": "toys",
            "автозапчасти": "automotive",
            "спорттовары": "sports",
            "другое": "other",
            "electronics": "electronics",
            "clothing": "clothing",
            "shoes": "shoes",
            "home_appliances": "home_appliances",
            "beauty": "beauty",
            "toys": "toys",
            "automotive": "automotive",
            "sports": "sports",
            "other": "other"
        }

        for row_num, row in enumerate(rows, start=2):
            if not row or not row[0]:  # Пропускаем пустые строки
                continue

            # Извлекаем данные
            track_code = str(row[0]).strip().upper() if row[0] else None
            product_name = str(row[1]).strip() if len(row) > 1 and row[1] else f"Товар {track_code}"
            category_raw = str(row[2]).strip().lower() if len(row) > 2 and row[2] else "other"
            quantity = row[3] if len(row) > 3 and row[3] else 1
            price = row[4] if len(row) > 4 and row[4] else 0.0
            weight = row[5] if len(row) > 5 and row[5] else 0.0
            country = str(row[6]).strip() if len(row) > 6 and row[6] else "China"

            # Валидация трек-кода
            if not track_code:
                result.add_error(f"Строка {row_num}", "Отсутствует трек-код")
                continue

            is_valid, error_msg = self.validator.validate_track_code(track_code)
            if not is_valid:
                result.add_error(track_code, f"Строка {row_num}: {error_msg}")
                continue

            # Проверка на дубликат
            existing = await self.product_repo.get_product_by_track_code(track_code)
            if existing:
                result.add_duplicate(track_code)
                continue

            # Определяем категорию
            category = category_map.get(category_raw, "other")

            # Валидация числовых полей
            try:
                quantity = int(quantity) if quantity else 1
                if quantity <= 0:
                    quantity = 1
            except (ValueError, TypeError):
                quantity = 1

            try:
                price = float(price) if price else 0.0
                if price < 0:
                    price = 0.0
            except (ValueError, TypeError):
                price = 0.0

            try:
                weight = float(weight) if weight else 0.0
                if weight < 0:
                    weight = 0.0
            except (ValueError, TypeError):
                weight = 0.0

            # Создаём товар
            try:
                total_value = quantity * price

                await self.product_repo.create_product(
                    track_code=track_code,
                    user_id=user_id,
                    product_name=product_name,
                    product_category=category,
                    country_from=country,
                    quantity=quantity,
                    unit_price_usd=price,
                    total_value_usd=total_value,
                    weight_kg=weight,
                    status=ProductStatus.CREATED
                )
                result.add_success(track_code)
                logger.info(f"Успешно импортирован трек-код из Excel: {track_code}")
            except Exception as e:
                error_msg = f"Строка {row_num}: Ошибка БД - {str(e)}"
                result.add_error(track_code, error_msg)
                logger.error(f"Ошибка при создании товара {track_code} из строки {row_num}: {e}")

        # Закрываем workbook
        wb.close()

        return result
